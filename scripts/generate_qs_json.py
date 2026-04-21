#!/usr/bin/env python3
"""
Generate QS World University Rankings 2026 JSON from Excel file.

This script parses the QS rankings Excel file and creates a JSON file
that can be used by the seed script.

Usage:
    python scripts/generate_qs_json.py
"""

import json
import sys
from pathlib import Path

try:
    from openpyxl import load_workbook
except ImportError:
    print("Error: openpyxl is required. Install with: pip install openpyxl")
    sys.exit(1)


def parse_rank_display(rank_str: str | None) -> tuple[str, int]:
    """Convert rank display string to display and numeric position.

    Examples:
        "1" -> ("1", 1)
        "701-710" -> ("701-710", 701)
        "1001-1200" -> ("1001-1200", 1001)
        "1401+" -> ("1401+", 1401)
        "31=" -> ("31=", 31)
    """
    if not rank_str:
        return ("", 0)

    rank_str = str(rank_str).strip()
    display = rank_str

    clean_rank = rank_str.replace("=", "").strip()

    if clean_rank.endswith("+"):
        return (display, int(clean_rank[:-1]))

    if "-" in clean_rank:
        return (display, int(clean_rank.split("-")[0]))

    try:
        return (display, int(clean_rank))
    except ValueError:
        return (display, 0)


def parse_score(score_val) -> float | None:
    """Parse a score value to float, handling empty values."""
    if score_val is None:
        return None
    if isinstance(score_val, (int, float)):
        return float(score_val)
    score_str = str(score_val).strip()
    if not score_str or score_str == "-" or score_str.lower() == "n/a":
        return None
    try:
        return float(score_str)
    except ValueError:
        return None


def find_header_row(sheet) -> int:
    """Find the row containing column headers."""
    for row_idx, row in enumerate(sheet.iter_rows(min_row=1, max_row=10), start=1):
        for cell in row:
            if cell.value and str(cell.value).lower() in ["rank", "2026 rank", "institution", "name"]:
                return row_idx
    return 1


def generate_qs_json():
    """Generate QS rankings JSON from Excel file."""
    script_dir = Path(__file__).parent.parent
    excel_file = script_dir / "2026 QS World University Rankings 1.3 (For qs.com).xlsx"

    if not excel_file.exists():
        print(f"Error: Excel file not found at {excel_file}")
        sys.exit(1)

    print(f"Loading Excel file: {excel_file}")
    workbook = load_workbook(excel_file, read_only=True, data_only=True)
    sheet = workbook.active

    header_row = find_header_row(sheet)
    headers = []
    for cell in sheet[header_row]:
        val = str(cell.value).lower().strip() if cell.value else ""
        headers.append(val)

    print(f"Found headers at row {header_row}: {headers[:15]}...")

    col_map = {}
    for idx, header in enumerate(headers):
        header_lower = header.lower()
        if header_lower == "2026" or ("2026" in header_lower and "rank" in header_lower):
            col_map["rank_2026"] = idx
        elif header_lower == "2025" or ("2025" in header_lower and "rank" in header_lower):
            col_map["rank_2025"] = idx
        elif "overall" in header_lower:
            col_map["overall_score"] = idx
        elif header_lower in ["institution", "institution name", "name"]:
            col_map["name"] = idx
        elif header_lower in ["location", "country", "location code"]:
            col_map["country"] = idx
        elif "academic" in header_lower and "reputation" in header_lower:
            col_map["academic_reputation"] = idx
        elif "employer" in header_lower and "reputation" in header_lower:
            col_map["employer_reputation"] = idx
        elif "faculty" in header_lower and "student" in header_lower:
            col_map["faculty_student_ratio"] = idx
        elif "citation" in header_lower:
            col_map["citations_per_faculty"] = idx
        elif "international" in header_lower and "faculty" in header_lower:
            col_map["international_faculty_ratio"] = idx
        elif "international" in header_lower and "student" in header_lower:
            col_map["international_students_ratio"] = idx
        elif "international" in header_lower and "research" in header_lower:
            col_map["international_research_network"] = idx
        elif "employment" in header_lower or "outcome" in header_lower:
            col_map["employment_outcomes"] = idx
        elif "sustainability" in header_lower:
            col_map["sustainability"] = idx
        elif header_lower in ["size", "institution size"]:
            col_map["size"] = idx
        elif header_lower in ["focus", "institution focus"]:
            col_map["focus"] = idx
        elif "research" in header_lower and "output" in header_lower:
            col_map["research_output"] = idx
        elif header_lower == "region":
            col_map["region"] = idx
        elif header_lower == "status":
            col_map["status"] = idx

    print(f"Column mapping: {col_map}")

    institutions = []
    data_start_row = header_row + 1

    for row in sheet.iter_rows(min_row=data_start_row):
        row_values = [cell.value for cell in row]

        if not any(row_values):
            continue

        rank_col = col_map.get("rank_2026", 0)
        name_col = col_map.get("name", 1)
        country_col = col_map.get("country", 2)

        rank_val = row_values[rank_col] if rank_col < len(row_values) else None
        name_val = row_values[name_col] if name_col < len(row_values) else None
        country_val = row_values[country_col] if country_col < len(row_values) else None

        if not name_val or not str(name_val).strip():
            continue

        rank_display, rank_position = parse_rank_display(rank_val)

        if rank_position == 0:
            continue

        overall_score = None
        if "overall_score" in col_map:
            overall_score = parse_score(row_values[col_map["overall_score"]])

        indicators = {}
        indicator_cols = [
            "academic_reputation",
            "employer_reputation",
            "faculty_student_ratio",
            "citations_per_faculty",
            "international_faculty_ratio",
            "international_students_ratio",
            "international_research_network",
            "employment_outcomes",
            "sustainability",
        ]
        for ind in indicator_cols:
            if ind in col_map:
                val = parse_score(row_values[col_map[ind]])
                if val is not None:
                    indicators[ind] = val

        institution = {
            "rank_display": rank_display,
            "rank_position": rank_position,
            "name": str(name_val).strip(),
            "country": str(country_val).strip() if country_val else "Unknown",
            "overall_score": overall_score,
        }

        if "rank_2025" in col_map:
            prev_rank = row_values[col_map["rank_2025"]]
            if prev_rank:
                institution["previous_rank_display"] = str(prev_rank).strip()

        if "size" in col_map:
            size_val = row_values[col_map["size"]]
            if size_val:
                institution["size"] = str(size_val).strip()

        if "focus" in col_map:
            focus_val = row_values[col_map["focus"]]
            if focus_val:
                institution["focus"] = str(focus_val).strip()

        if "research_output" in col_map:
            ro_val = row_values[col_map["research_output"]]
            if ro_val:
                institution["research_output"] = str(ro_val).strip()

        if "region" in col_map:
            region_val = row_values[col_map["region"]]
            if region_val:
                institution["region"] = str(region_val).strip()

        if "status" in col_map:
            status_val = row_values[col_map["status"]]
            if status_val:
                institution["status"] = str(status_val).strip()

        if indicators:
            institution["indicators"] = indicators

        institutions.append(institution)

    workbook.close()

    institutions.sort(key=lambda x: x["rank_position"])

    output = {
        "ranking_source": "qs_world",
        "ranking_year": 2026,
        "source_url": "https://www.topuniversities.com/world-university-rankings/2026",
        "institutions": institutions,
    }

    output_dir = script_dir / "data"
    output_dir.mkdir(exist_ok=True)

    output_file = output_dir / "qs_world_rankings_2026.json"
    with open(output_file, "w", encoding="utf-8") as f:
        json.dump(output, f, indent=2, ensure_ascii=False)

    print(f"\nGenerated {len(institutions)} institutions to {output_file}")

    if institutions:
        print(f"\nSample entry (rank #{institutions[0]['rank_position']}):")
        print(json.dumps(institutions[0], indent=2))


if __name__ == "__main__":
    generate_qs_json()
